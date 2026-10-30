import csv
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from eda_toolkit.core_io import list_templates, load_recordings, read_template
from eda_toolkit.ui.workbook import export_workbook, import_workbook


def _fill_examples(path, *, skip=()):
    workbook_path = export_workbook(path)
    workbook = load_workbook(workbook_path)
    for template in list_templates():
        name = template["name"]
        if name in skip:
            continue
        frame = read_template(name, "example")
        sheet = workbook[template["display_name"]]
        for row_number, values in enumerate(frame.itertuples(index=False, name=None), start=2):
            for column_number, value in enumerate(values, start=1):
                if value is not None and value != "":
                    sheet.cell(row_number, column_number, value)
    workbook.save(workbook_path)
    return workbook_path


def test_export_workbook_has_chinese_sheets_canonical_headers_and_no_example_rows(tmp_path):
    path = export_workbook(tmp_path)
    workbook = load_workbook(path, data_only=False)

    assert path.name == "EDA研究工作簿.xlsx"
    assert {"填写说明", "完整性状态", "字段字典"}.issubset(workbook.sheetnames)
    for template in list_templates():
        sheet = workbook[template["display_name"]]
        assert [cell.value for cell in sheet[1]] == list(read_template(template["name"], "template").columns)
        assert sheet.max_row == 1
        assert sheet[1][0].comment is not None


def test_example_workbook_import_snapshots_all_valid_tables_and_loads_record(tmp_path):
    workbook_path = _fill_examples(tmp_path / "examples.xlsx", skip=("input_manifest",))
    workbook = load_workbook(workbook_path)
    # Verify a missing signal is exported as an empty CSV field, never the text "nan".
    workbook["信号"][4][5].value = None
    workbook.save(workbook_path)

    result = import_workbook(workbook_path, tmp_path / "imports")

    assert result["requirements"][0]["status"] == "complete"
    assert result["statuses"]["signals"]["status"] == "complete"
    assert result["statuses"]["statistical_design"]["status"] == "complete"
    assert result["statuses"]["input_manifest"]["status"] == "optional_empty"
    assert result["statistics_designs"][0]["design_type"] == "paired"
    assert result["statistics_designs"][0]["subject_column"] == "subject_id"
    assert Path(result["bindings"]["signals"][0]).is_file()

    with Path(result["snapshots"]["signals"]).open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert rows[-1]["eda_uS"] == ""
    assert rows[-1]["subject_id"] == "SYNTHETIC_SUBJECT_001"

    records, tables, _ = load_recordings(options={"input": result["bindings"], "mapping": {"time_unit": "s"}})
    assert len(records) == 1
    assert records[0]["metadata"]["recording_id"] == "SYNTHETIC_REC_001"
    assert records[0]["metadata"]["subject_id"] == "SYNTHETIC_SUBJECT_001"
    assert records[0]["metadata"]["channel"] == "EDA"
    assert records[0]["metadata"]["unit"] == "uS"
    assert len(tables["events"]) == 1


def test_missing_required_field_is_reported_and_invalid_sheet_is_not_bound(tmp_path):
    workbook_path = _fill_examples(tmp_path / "missing.xlsx", skip=("input_manifest",))
    workbook = load_workbook(workbook_path)
    # Remove event_id from the canonical header and its values.
    workbook["事件"].delete_cols(2)
    workbook.save(workbook_path)

    result = import_workbook(workbook_path, tmp_path / "imports")
    event_status = result["statuses"]["events"]

    assert event_status["status"] == "invalid"
    assert "event_id" in event_status["missing_fields"]
    assert any(error["code"] == "required_column" for error in event_status["errors"])
    assert result["bindings"]["events"] is None
    assert result["statuses"]["participants"]["status"] == "complete"


def test_empty_optional_sheets_and_config_required_statuses(tmp_path):
    workbook_path = export_workbook(tmp_path / "empty.xlsx")
    result = import_workbook(
        workbook_path,
        tmp_path / "imports",
        config={"statistics": {"enabled": True}, "synchrony": {"enabled": True}},
    )

    assert result["statuses"]["events"]["status"] == "optional_empty"
    assert result["statuses"]["statistical_design"]["status"] == "missing_required"
    assert result["statuses"]["group_membership"]["status"] == "missing_required"
    assert result["statuses"]["sync_anchors"]["status"] == "missing_required"
    assert result["statuses"]["participants"]["status"] == "optional_empty"
    assert result["statistics_designs"] == []
    assert result["bindings"]["signals"] == []


def test_relative_manifest_path_resolves_from_workbook_and_load_recordings_uses_bindings(tmp_path):
    raw_signal = tmp_path / "raw_signal.csv"
    raw_signal.write_text(
        "time_s,eda_uS\n0,5.0\n1,6.0\n",
        encoding="utf-8",
    )
    workbook_path = _fill_examples(
        tmp_path / "manifest.xlsx",
        skip=("signals", "input_manifest", "events", "conditions", "group_membership", "sync_anchors", "manual_qc", "statistical_design"),
    )
    workbook = load_workbook(workbook_path)
    manifest = workbook["输入清单"]
    sample = read_template("input_manifest", "example").iloc[0].to_dict()
    sample["file_path"] = "raw_signal.csv"
    for column_number, name in enumerate(read_template("input_manifest", "template").columns, start=1):
        value = sample.get(name)
        if value is not None and value != "":
            manifest.cell(2, column_number, value)
    workbook.save(workbook_path)

    result = import_workbook(workbook_path, tmp_path / "imports")
    snapshot_manifest = Path(result["snapshots"]["input_manifest"])
    with snapshot_manifest.open(encoding="utf-8-sig", newline="") as stream:
        imported_path = next(csv.DictReader(stream))["file_path"]
    assert Path(imported_path) == raw_signal.resolve()
    assert result["bindings"]["manifest"] == str(snapshot_manifest)

    records, _, _ = load_recordings(options={"input": result["bindings"], "mapping": {}})
    metadata = records[0]["metadata"]
    assert metadata["recording_id"] == "SYNTHETIC_REC_001"
    assert metadata["subject_id"] == "SYNTHETIC_SUBJECT_001"
    assert metadata["channel"] == "EDA"
    assert metadata["unit"] == "uS"
    assert Path(metadata["source_path"]) == raw_signal.resolve()

